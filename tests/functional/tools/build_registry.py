"""Generate registry/ecs_test_registry.json from the functional-test workbook (the source of truth).

Usage (repo root):  python tests/functional/tools/build_registry.py [path/to/ECS_Consolidated_Functional_Test_Pack_UC01_UC20.xlsx]

Preserves UC ID, Functional Test ID, description, preconditions, steps, expected result, priority and test type VERBATIM.
Adds only derived fields: topic, automation_capability (family), dependencies, markers.
"""

from __future__ import annotations

import json
import os
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
FUNC = HERE.parent
sys.path.insert(0, str(FUNC.parent))      # tests/ -> package "functional"
sys.path.insert(0, str(FUNC.parents[1]))  # repo root -> ecs imports

import openpyxl  # noqa: E402

from functional.framework.classify import FAMILY_DEPENDENCIES, FAMILY_MARKERS, core_family, parse_topic  # noqa: E402

DEFAULT_XLSX = os.environ.get("ECS_FT_WORKBOOK", str(Path.home() / "Downloads" / "ECS_Consolidated_Functional_Test_Pack_UC01_UC20.xlsx"))
OUT = FUNC / "registry" / "ecs_test_registry.json"
TYPE_MARKERS = {"Negative": ["negative"], "Positive": ["positive"], "Functional": []}


def build(xlsx: str) -> dict:
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    names = {r[0]: r[1] for r in wb["SUMMARY"].iter_rows(min_row=2, values_only=True) if r[0]}
    registry: dict = {"source_workbook": Path(xlsx).name, "use_cases": {}}
    for ws in wb:
        if not (ws.title.startswith("UC") and ws.title[2:].isdigit()):
            continue
        tests = []
        for row in ws.iter_rows(min_row=5, values_only=True):
            if not (row[0] and str(row[0]).startswith(ws.title + "-FT")):
                continue
            tid, name, pre, steps, expected, priority, ttype = (str(c).strip() if c is not None else "" for c in row[:7])
            slug, family = parse_topic(name)
            family = family or core_family(name)
            markers = ["functional", ws.title.lower(), *FAMILY_MARKERS.get(family, ["api"]), *TYPE_MARKERS.get(ttype, [])]
            tests.append({"id": tid, "name": name, "preconditions": pre, "steps": steps, "expected": expected,
                          "priority": priority, "type": ttype, "topic": slug, "automation_capability": family,
                          "dependencies": FAMILY_DEPENDENCIES.get(family, ["api_client"]),
                          "markers": list(dict.fromkeys(markers))})
        registry["use_cases"][ws.title] = {"name": names.get(ws.title, ws.title), "tests": tests}
    return registry


if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    reg = build(src)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(reg, indent=2, ensure_ascii=False), encoding="utf-8")
    total = sum(len(u["tests"]) for u in reg["use_cases"].values())
    print(f"wrote {OUT} - {len(reg['use_cases'])} use cases, {total} functional tests")
